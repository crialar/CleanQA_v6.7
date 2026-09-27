"""Explicit exports in RAM. Incremental XML writer has no whole-file size cutoff."""
import csv,io,zipfile
from lxml import etree as E
from document_core import local_name,tmx_variant,XML_LANG,XLIFF_URI
import qa_checker as q


def report_results(parent):
    state=parent['state'];files=state.get('files',[]);categories={}
    for item in parent['issues']:
        cid=item['category_id']
        if cid not in categories:categories[cid]=dict(q.CATEGORY_METADATA[cid],issues=[],count=0)
        issue=dict(item)
        if item.get('_issue_id') in parent.get('ignored',set()):issue['note']='Ignored in the viewer. '+issue.get('note','')
        categories[cid]['issues'].append(issue);categories[cid]['count']+=1
    notes=list(state.get('notices',[]))
    notes.insert(0,'Status: '+state['state'])
    if state.get('error'):notes.append(state['error'])
    if parent['edits']:notes.append('There are edits that have not been reanalyzed. Findings reflect the previous analysis.')
    first=files[0] if files else {}
    return dict(files=files,filename=', '.join(f['name'] for f in files),format=first.get('format',''),monolingual=first.get('format')=='docx',source_lang=first.get('source_lang',''),target_lang=first.get('target_lang',''),notices=notes,coverage={'complete':state.get('complete',False)},segment_count=state.get('segment_count',0),categories=categories,segments_data=parent.get('segments',[]),summary={'total':len(parent['issues']),'high':0,'low':0})


class XMLTarget:
    def __init__(self,writer,fmt,target,overrides,pulse):
        self.writer=writer;self.fmt=fmt;self.target=target;self.overrides=overrides;self.pulse=pulse
        self.contexts=[];self.ns=[];self.pending={};self.nodes=[];self.unit=None;self.index=0
    def start_ns(self,prefix,uri):self.pending[prefix or None]=uri
    def end_ns(self,prefix):pass
    def start(self,tag,attrs):
        mapping=dict(self.ns[-1]) if self.ns else {};mapping.update(self.pending);declared=self.pending;self.pending={};self.ns.append(mapping)
        if self.unit is not None or local_name(E.Element(tag))==('tu' if self.fmt=='tmx' else 'trans-unit'):
            el=E.Element(tag,attrib=attrs,nsmap=mapping or None)
            if self.nodes:self.nodes[-1].append(el)
            else:self.unit=el
            self.nodes.append(el)
        else:
            context=self.writer.element(tag,attrib=attrs,nsmap=declared or None);context.__enter__();self.contexts.append(context)
    def data(self,text):
        if self.nodes:
            el=self.nodes[-1]
            if len(el):el[-1].tail=(el[-1].tail or '')+text
            else:el.text=(el.text or '')+text
        else:self.writer.write(text)
    def end(self,tag):
        if self.nodes:
            self.nodes.pop()
            if not self.nodes:
                self.index+=1;edit=self.overrides.get(str(self.index))
                if edit is not None:
                    if self.fmt=='tmx':
                        target=tmx_variant(self.unit,self.target)
                        if target is None:
                            tuv=E.SubElement(self.unit,'tuv');tuv.set(XML_LANG,self.target);target=E.SubElement(tuv,'seg')
                    else:
                        target=self.unit.find('{'+XLIFF_URI+'}target')
                        if target is None:
                            target=E.Element('{'+XLIFF_URI+'}target')
                            source=self.unit.find('{'+XLIFF_URI+'}source')
                            self.unit.insert(list(self.unit).index(source)+1 if source is not None else 0,target)
                    q._replace_target_with_plain_text(target,edit)
                self.writer.write(self.unit);self.unit=None
                if self.index%100==0:self.pulse('Writing XML',self.index,None)
        else:self.contexts.pop().__exit__(None,None,None)
        self.ns.pop()
    def comment(self,text):
        el=E.Comment(text)
        if self.nodes:self.nodes[-1].append(el)
        else:self.writer.write(el)
    def pi(self,target,text):
        el=E.ProcessingInstruction(target,text)
        if self.nodes:self.nodes[-1].append(el)
        else:self.writer.write(el)
    def doctype(self,name,pubid,system):pass  # External DTD is not fetched or required in output.
    def close(self):return self.index


def corrected_xml(source,dest,fmt,target,overrides,pulse):
    from pathlib import Path
    stream=source.open('rb') if isinstance(source,Path) else source
    try:
        if hasattr(stream,'seek'):stream.seek(0)
        with E.xmlfile(dest,encoding='UTF-8') as writer:
            writer.write_declaration()
            handler=XMLTarget(writer,fmt,target,overrides,pulse)
            parser=E.XMLParser(target=handler,resolve_entities=False,load_dtd=False,no_network=True,huge_tree=False,recover=False)
            while True:
                block=stream.read(65536)
                if not block:break
                parser.feed(block)
            parser.close()
    finally:
        if isinstance(source,Path):stream.close()


def export_job(kind,parent,pulse):
    if kind in {'html','csv','xlsx'}:
        pulse('Preparing interactive report' if kind=='html' else ('Preparing Excel' if kind=='xlsx' else 'Preparing CSV'),0,None)
        result=report_results(parent)
        data=q.export_qa_report(result,'csv' if kind=='xlsx' else kind,target_overrides={str(k):v for k,v in parent['edits'].items()})
        if kind in {'csv','xlsx'}:
            # Keep a rectangular CSV and expose coverage even with zero findings.
            rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
            header=rows[0]
            state=parent['state']
            metadata=[state['state'], 'true' if state.get('complete',False) else 'false',
                      '\n'.join(result['notices']),result['target_lang']]
            from document_core import safe_cell
            metadata=[safe_cell(v) for v in metadata]
            out=io.StringIO();writer=csv.writer(out,quoting=csv.QUOTE_ALL)
            writer.writerow(header+['Analysis state','Analysis complete','Coverage notices','Target language'])
            for row in rows[1:] or [['']*len(header)]:
                writer.writerow(row+metadata)
            data=out.getvalue().encode('utf-8-sig')
        if kind=='xlsx':data=_csv_to_xlsx(data,pulse)
        return data,'QA_report.'+kind
    if kind!='edited':raise ValueError('Unknown export format.')
    from qa_jobs import _source
    out=io.BytesIO();edits=dict(parent['overrides']);edits.update({str(k):v for k,v in parent['edits'].items()})
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED,allowZip64=True) as z:
        for f,entry in zip(parent['state']['files'],parent['files']):
            if f['format'] not in {'tmx','mqxliff'}:raise ValueError('Only bilingual XML can be edited.')
            local={str(int(s)-f['offset']):v for s,v in edits.items() if f['offset']<int(s)<=f['offset']+f['segment_count']}
            with z.open('Corrected_'+f['name'],'w',force_zip64=True) as output:
                corrected_xml(_source(entry),output,f['format'],f['target_lang'],local,pulse)
    return out.getvalue(),'QA_corrected.zip'


EXCEL_MAX_ROWS = 1048576


def _csv_to_xlsx(data, pulse):
    """Build the report using the application's existing Excel dependency, in RAM.

    Normal workbook mode deliberately avoids write-only temporary disk files.
    Cells are text: translated content can never become executable formulas.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    rows = csv.reader(io.StringIO(data.decode('utf-8-sig')))
    header = next(rows)
    book = Workbook()
    book.remove(book.active)
    sheet = None
    fill = PatternFill('solid', fgColor='174A70')
    font = Font(color='FFFFFF', bold=True)
    wrap = Alignment(vertical='top', wrap_text=True)
    widths = [35,12,30,22,12,65,45,65,65,65,18,20,60,20]

    # Explicit row indices avoid implicit conversions and preserve empty cells.
    def write_row(index, values):
        for col, value in enumerate(values, 1):
            if len(value) > 32767:
                raise ValueError('Excel cannot store a report cell longer than 32,767 characters. Use the interactive HTML report; no truncated Excel was generated.')
            cell = sheet.cell(index,col,value)
            cell.data_type = 's'
            if index > 1 and header[col-1] == 'Segment' and value.isascii() and value.isdigit() and len(value) <= 15 and str(int(value)) == value:
                cell.value = int(value)
            cell.alignment = wrap

    for count, values in enumerate(rows, 1):
        if sheet is None or sheet.max_row >= EXCEL_MAX_ROWS:
            if sheet is not None:sheet.auto_filter.ref = sheet.dimensions
            sheet = book.create_sheet('QA report' if not book.worksheets else f'QA report {len(book.worksheets)+1}')
            write_row(1,header)
            sheet.freeze_panes = 'C2'
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 32
            for cell in sheet[1]:cell.fill=fill;cell.font=font
            for col in range(1,len(header)+1):
                sheet.column_dimensions[get_column_letter(col)].width = widths[col-1] if col<=len(widths) else 24
        write_row(sheet.max_row+1,values)
        if count%250==0:pulse('Preparing Excel',count,None)
    if sheet is None:
        sheet=book.create_sheet('QA report');write_row(1,header);sheet.freeze_panes='C2'
    sheet.auto_filter.ref=sheet.dimensions
    output=io.BytesIO()
    book.save(output)
    book.close()
    pulse('Excel ready',count if 'count' in locals() else 0,None)
    return output.getvalue()
