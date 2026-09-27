"""User-defined confusable groups; no external services or built-in linguistic claims."""
import io
import re
from document_core import checked_zip


def _groups(rows):
    groups, seen = [], set()
    for row in rows:
        members, keys = [], set()
        for value in row:
            if value is None:
                continue
            text = str(value).strip()
            key = text.casefold()
            if text and key not in keys:
                if len(text) > 500:
                    raise ValueError('A confusable member exceeds 500 characters.')
                members.append(text)
                keys.add(key)
        signature = tuple(sorted(keys))
        if len(members) >= 2 and signature not in seen:
            groups.append(tuple(members))
            seen.add(signature)
        if len(groups) > 10000:
            raise ValueError('Confusable list exceeds 10,000 groups.')
    return groups


def parse_custom_pairs(text):
    if len(text) > 2_000_000:
        raise ValueError('Confusable text exceeds the size limit.')
    return _groups(re.split(r'[|/]', line) for line in text.splitlines()
                   if line.strip() and not line.lstrip().startswith('#'))


def parse_pairs_xlsx(data, filename='pairs.xlsx'):
    from openpyxl import load_workbook
    with checked_zip(data):
        pass
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            return _groups(row for sheet in wb for row in sheet.iter_rows(values_only=True))
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f'Could not read confusable pairs from {filename}: {exc}') from exc
